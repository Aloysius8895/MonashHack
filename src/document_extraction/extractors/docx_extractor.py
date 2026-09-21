"""Extractor for Word SI/BL attachments.

Format: a single table, 2 columns (label, value). Labels in this dataset
are bilingual, e.g. "Shipper (Principal or Seller) (发货人)" -
normalize_label() strips the CJK part before alias/keyword matching.
"""
import io

import docx

from ..fields import classify_label
from ..normalize import normalize_label


def extract(data):
    document = docx.Document(io.BytesIO(data))
    hits = {}
    text_parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table_idx, table in enumerate(document.tables):
        for row_idx, row in enumerate(table.rows, start=1):
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) < 2 or not cells[0]:
                continue
            label, value = cells[0], cells[1]
            text_parts.append(f"{label} {value}")
            field_name = classify_label(normalize_label(label))
            if field_name and field_name not in hits:
                hits[field_name] = (label, value, f"table {table_idx} row {row_idx}")
    return hits, "\n".join(text_parts), {}
